import sys, unittest
from pathlib import Path
from uuid import uuid4
from docx import Document
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from resume_base import inspect_docx


class ResumeBaseTests(unittest.TestCase):
    def test_safe_groups(self):
        folder = Path('output/test-runs') / str(uuid4())
        folder.mkdir(parents=True)
        doc = Document()
        doc.add_paragraph('Synthetic Applicant Identity')
        doc.add_paragraph('EDUCATION')
        doc.add_paragraph('Synthetic academic achievement that must stay unchanged', style='List Bullet')
        doc.add_paragraph('PROJECTS')
        doc.add_paragraph('Project One')
        doc.add_paragraph('Built JavaScript user interfaces for the first project.', style='List Bullet')
        doc.add_paragraph('Built Python SQL reporting for the first project.', style='List Bullet')
        doc.add_paragraph('Project Two')
        doc.add_paragraph('Built Java systems for a separate second project.', style='List Bullet')
        p = doc.add_paragraph(style='List Bullet')
        p.add_run('Mixed bold skill ').bold = True
        p.add_run('should remain exactly where it was written.')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Table content must remain unchanged for format safety.'
        path = folder / 'source.docx'
        doc.save(path)
        original = path.read_bytes()
        result = inspect_docx(path)
        self.assertEqual(result['groups'], [[5, 6]])
        self.assertEqual([s['paragraph'] for s in result['slots']], [5, 6, 8])
        self.assertIn('Synthetic Applicant Identity', result['fixedText'])
        self.assertIn('Synthetic academic achievement that must stay unchanged', result['fixedText'])
        self.assertEqual(original, path.read_bytes())


if __name__ == '__main__':
    unittest.main()
